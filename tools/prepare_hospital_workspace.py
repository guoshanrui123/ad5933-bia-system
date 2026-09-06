"""Create a local empty data workspace; never copy clinical records."""

from pathlib import Path
import shutil

from openpyxl import Workbook
import hospital_capture_app as app
import hospital_calibration_core as calibration


def main():
    root = app.application_root()
    reference = root / "calibration_reference"
    reference.mkdir(parents=True, exist_ok=True)
    source = Path(__file__).resolve().parents[1] / "calibration" / "reference_loads"
    for _, filename, *_ in calibration.CALIBRATION_LOADS:
        target = reference / filename
        if not target.exists():
            shutil.copy2(source / filename, target)
    for name in ("raw_txt", "qc_results", "pending_excel_records"):
        (root / name).mkdir(exist_ok=True)
    workbook_path = root / app.WORKBOOK_NAME
    if not workbook_path.exists():
        workbook = Workbook()
        for index, (name, headers) in enumerate((
            ("Sessions", app.SESSION_HEADERS),
            ("Frequency_Summary", app.FREQUENCY_HEADERS),
            ("Sweep_Level", app.SWEEP_HEADERS),
        )):
            sheet = workbook.active if index == 0 else workbook.create_sheet()
            sheet.title = name
            sheet.cell(1, 1, "Empty research template; no clinical validation implied")
            for column, header in enumerate(headers, 1):
                sheet.cell(4, column, header)
            sheet.freeze_panes = "A5"
        with workbook_path.open("xb") as stream:
            workbook.save(stream)
    calibration.build_engine(reference)
    print(f"Prepared local workspace: {root}")


if __name__ == "__main__":
    main()
